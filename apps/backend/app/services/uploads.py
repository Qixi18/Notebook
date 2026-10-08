"""Safe multi-format ingestion; one committed material and queued job per upload."""

from __future__ import annotations

import hashlib
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import BadZipFile, ZipFile

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.models import Material, ProcessingJob, new_id
from app.notes.notebook import bump_order_revision


class UploadRejected(ValueError):
    def __init__(self, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.status_code = status_code


def _validate_container(path: Path, extension: str) -> None:
    if extension == ".pdf":
        with path.open("rb") as source:
            if source.read(5) != b"%PDF-":
                raise UploadRejected("文件扩展名是 PDF，但文件签名无效", 415)
        try:
            from pypdf import PdfReader
            page_count = len(PdfReader(str(path)).pages)
        except Exception as exc:
            raise UploadRejected("PDF 文件损坏或无法读取", 415) from exc
        if page_count > settings.max_pages:
            raise UploadRejected(f"PDF 页面数超过本地限制（最多 {settings.max_pages} 页）", 413)
        return
    try:
        with ZipFile(path) as archive:
            names = set(archive.namelist())
            required = {"ppt/presentation.xml"} if extension == ".pptx" else {"word/document.xml", "[Content_Types].xml"}
            if not required.issubset(names) or archive.testzip() is not None:
                raise UploadRejected(f"文件不是可读取的 {extension[1:].upper()}", 415)
        if extension == ".pptx":
            from pptx import Presentation

            page_count = len(Presentation(str(path)).slides)
            if page_count > settings.max_pages:
                raise UploadRejected(f"PPTX 页面数超过本地限制（最多 {settings.max_pages} 页）", 413)
        elif extension == ".docx":
            from docx import Document

            document = Document(str(path))
            position_count = len(document.paragraphs) + len(document.tables)
            if position_count > settings.max_pages:
                raise UploadRejected(f"DOCX 章节/位置数超过本地限制（最多 {settings.max_pages} 个）", 413)
    except UploadRejected:
        raise
    except BadZipFile as exc:
        raise UploadRejected(f"文件不是可读取的 {extension[1:].upper()}", 415) from exc
    except Exception as exc:
        raise UploadRejected(f"文件不是可读取的 {extension[1:].upper()}", 415) from exc


def ingest_document(
    db: Session, *, course_id: str, source: object, filename: str,
    lecture_title: str, topic_title: str, media_type: str | None,
    allow_duplicate: bool = False, idempotency_key: str | None = None,
) -> tuple[Material, ProcessingJob]:
    extension = Path(filename).suffix.lower()
    if extension not in settings.allowed_extensions:
        raise UploadRejected("当前支持 PPTX、PDF 和 DOCX 文件", 415)
    key = hashlib.sha256(f"{course_id}:{idempotency_key}".encode()).hexdigest() if idempotency_key else None
    if key:
        prior = db.scalar(select(ProcessingJob).where(ProcessingJob.idempotency_key == key))
        if prior is not None:
            if prior.material.course_id != course_id:
                raise UploadRejected("重复请求不属于当前课程", 409)
            return prior.material, prior

    settings.originals_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256()
    size = 0
    temporary = None
    destination = None
    try:
        with NamedTemporaryFile(dir=settings.originals_dir, prefix="upload-", suffix=".tmp", delete=False) as output:
            temporary = Path(output.name)
            while chunk := source.read(1024 * 1024):
                size += len(chunk)
                if size > settings.max_upload_bytes:
                    raise UploadRejected("文件超过本地配置的大小限制", 413)
                output.write(chunk)
                digest.update(chunk)
        if size == 0:
            raise UploadRejected("文件为空或无法读取", 415)
        _validate_container(temporary, extension)
        content_hash = digest.hexdigest()
        duplicate = db.scalar(select(Material).where(
            Material.course_id == course_id, Material.content_hash == content_hash,
            Material.deleted_at.is_(None),
        ))
        if duplicate is not None and not allow_duplicate:
            raise UploadRejected(f"该课程已上传相同文件：{duplicate.lecture_title}。如需作为新讲次导入，请确认。", 409)
        material_id = new_id()
        # Serialize chapter allocation with reorders and other uploads.
        bump_order_revision(db, course_id)
        last_order = db.scalar(select(func.max(Material.chapter_order)).where(Material.course_id == course_id))
        destination = settings.originals_dir / f"{material_id}{extension}"
        os.replace(temporary, destination)
        temporary = None
        material = Material(
            id=material_id, course_id=course_id,
            lecture_title=lecture_title.strip()[:200] or filename[:200],
            topic_title=topic_title.strip()[:200] or None,
            original_filename=Path(filename).name[:255], stored_filename=destination.name,
            media_type=media_type, size_bytes=size, content_hash=content_hash,
            status="pending",
            chapter_order=0 if last_order is None else last_order + 1,
        )
        job = ProcessingJob(material_id=material_id, kind="parse_material", idempotency_key=key)
        db.add_all([material, job])
        db.commit()
        db.refresh(material)
        db.refresh(job)
        return material, job
    except Exception:
        db.rollback()
        if destination is not None:
            destination.unlink(missing_ok=True)
        raise
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def ingest_pptx(db: Session, **kwargs):
    """Backward-compatible name retained for integrations from phase one."""
    return ingest_document(db, **kwargs)
