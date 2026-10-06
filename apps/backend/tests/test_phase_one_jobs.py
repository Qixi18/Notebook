from __future__ import annotations

from dataclasses import replace
from io import BytesIO

import pytest
from pptx import Presentation
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker

from app.core.config import settings
from app.db.database import Base
from app.db.models import Course, Material, ProcessingJob
from app.services import uploads
from app.services.uploads import UploadRejected, ingest_pptx
from app.workers import runner


def _pptx() -> BytesIO:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "第一讲"
    result = BytesIO()
    presentation.save(result)
    result.seek(0)
    return result


@pytest.fixture
def database(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{(tmp_path / 'test.sqlite3').as_posix()}")

    @event.listens_for(engine, "connect")
    def foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    monkeypatch.setattr(uploads, "settings", replace(settings, data_dir=tmp_path))
    monkeypatch.setattr(runner, "SessionLocal", factory)
    with factory() as db:
        course = Course(name="课程 A")
        db.add(course)
        db.commit()
        yield factory, course.id
    engine.dispose()


def test_upload_is_queued_and_idempotent(database):
    factory, course_id = database
    with factory() as db:
        first, job = ingest_pptx(
            db, course_id=course_id, source=_pptx(), filename="lecture.pptx",
            lecture_title="第 1 讲", topic_title="", media_type=None,
            idempotency_key="request-1",
        )
        assert job.status == "pending"
        same, same_job = ingest_pptx(
            db, course_id=course_id, source=_pptx(), filename="lecture.pptx",
            lecture_title="第 1 讲", topic_title="", media_type=None,
            idempotency_key="request-1",
        )
        assert (same.id, same_job.id) == (first.id, job.id)
        with pytest.raises(UploadRejected) as duplicate:
            ingest_pptx(
                db, course_id=course_id, source=_pptx(), filename="lecture.pptx",
                lecture_title="第 2 讲", topic_title="", media_type=None,
            )
        assert duplicate.value.status_code == 409
        assert len(db.scalars(select(Material)).all()) == 1
        assert runner.claim_next_job() == job.id
        assert runner.claim_next_job() is None
        db.refresh(job)
        assert job.status == "processing" and job.attempt_count == 1


def test_interrupted_job_is_retryable_without_duplicate_rows(database):
    factory, course_id = database
    with factory() as db:
        material, job = ingest_pptx(
            db, course_id=course_id, source=_pptx(), filename="lecture.pptx",
            lecture_title="第 1 讲", topic_title="", media_type=None,
        )
        assert runner.claim_next_job() == job.id
        db.refresh(job)
        job.heartbeat_at = None
        db.commit()
        assert runner.mark_interrupted_jobs() == 1
        db.refresh(job)
        db.refresh(material)
        assert job.status == "failed" and job.error_code == "interrupted"
        assert material.status == "failed"
        retry = ProcessingJob(material_id=material.id, kind="parse_material")
        material.status = "pending"
        db.add(retry)
        db.commit()
        assert runner.claim_next_job() == retry.id


def test_upload_rejects_disguised_file_and_cleans_temporary(database, tmp_path):
    factory, course_id = database
    with factory() as db, pytest.raises(UploadRejected) as rejected:
        ingest_pptx(
            db, course_id=course_id, source=BytesIO(b"not a pptx"),
            filename="fake.pptx", lecture_title="bad", topic_title="", media_type=None,
        )
    assert rejected.value.status_code == 415
    assert list((tmp_path / "originals").iterdir()) == []
