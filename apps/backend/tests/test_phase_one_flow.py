"""Exercise the real API, migration, queue and parser in an isolated data directory."""

from __future__ import annotations

import os
import subprocess
import sys
import textwrap


def test_isolated_pptx_flow(tmp_path):
    code = textwrap.dedent("""
        from io import BytesIO
        from fastapi.testclient import TestClient
        from pptx import Presentation
        from app.main import app
        from app.workers.runner import run_once

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = '测试知识点'
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = '核心概念与学习方法'
        file = BytesIO()
        presentation.save(file)
        content = file.getvalue()

        with TestClient(app) as client:
            course = client.post('/api/v1/courses', json={'name':'阶段一测试课程'}).json()
            course_id = course['id']
            upload = client.post(f'/api/v1/courses/{course_id}/materials',
                files={'file': ('one.pptx', content, 'application/vnd.openxmlformats-officedocument.presentationml.presentation')},
                data={'lecture_title':'第一讲'}, headers={'Idempotency-Key':'first'})
            assert upload.status_code == 201, upload.text
            first = upload.json()
            assert first['job']['status'] == 'pending'
            assert run_once()
            assert client.get(f"/api/v1/jobs/{first['job']['id']}").json()['status'] == 'completed'
            pages = client.get(f"/api/v1/materials/{first['material']['id']}/pages").json()
            assert len(pages) == 1 and '测试知识点' in pages[0]['raw_text']
            notes = client.get(f'/api/v1/courses/{course_id}/notes').json()
            assert notes
            note = notes[0]
            edited = client.patch(f"/api/v1/notes/{note['id']}", json={
                'content_markdown':'# 我的用户正文',
                'expected_revision_number':note['revision_number'],
            })
            assert edited.status_code == 200, edited.text
            second = client.post(f'/api/v1/courses/{course_id}/materials',
                files={'file': ('two.pptx', content, 'application/vnd.openxmlformats-officedocument.presentationml.presentation')},
                data={'lecture_title':'第二讲', 'allow_duplicate':'true'})
            assert second.status_code == 201, second.text
            assert run_once()
            protected = client.get(f"/api/v1/notes/{note['id']}").json()
            assert protected['content_markdown'] == '# 我的用户正文'
            assert protected['user_locked']
            first_id = first['material']['id']
            preview = client.get(f'/api/v1/materials/{first_id}/deletion-preview').json()
            assert preview['pages'] == 1 and preview['user_notes_protected'] >= 1
            assert client.delete(f'/api/v1/materials/{first_id}').status_code == 200
            assert client.get(f'/api/v1/materials/{first_id}/pages').status_code == 404
            assert client.post(f'/api/v1/materials/{first_id}/restore').status_code == 200
            assert client.get(f'/api/v1/materials/{first_id}/pages').status_code == 200
            assert client.get(f"/api/v1/notes/{note['id']}").json()['content_markdown'] == '# 我的用户正文'
        print('phase-one-flow-ok')
    """)
    env = os.environ.copy()
    env.update({
        "NOTEBOOK_DATA_DIR": str(tmp_path),
        "DEEPSEEK_API_KEY": "", "TAVILY_API_KEY": "",
        "EMBEDDING_API_KEY": "", "EMBEDDING_BASE_URL": "", "EMBEDDING_MODEL": "",
    })
    completed = subprocess.run(
        [sys.executable, "-c", code], capture_output=True, text=True,
        env=env, timeout=90, check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr
    assert "phase-one-flow-ok" in completed.stdout
