"""Format parser and evidence coverage regression tests."""

from __future__ import annotations

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfWriter

from app.main import app
from app.parsers.docx_parser import parse_docx
from app.parsers.pdf_parser import parse_pdf
from app.parsers.pptx_parser import parse_pptx_document


def test_pptx_tables_have_stable_objects(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table = slide.shapes.add_table(2, 2, 0, 0, 1000000, 1000000).table
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "状态"
    table.cell(1, 1).text = "完成"
    path = tmp_path / "table.pptx"
    presentation.save(path)
    document = parse_pptx_document(path)
    assert document.format == "pptx"
    assert document.pages[0].stable_location_key == "slide:1"
    assert document.pages[0].blocks[0].extraction_method == "table_parse"
    assert "字段" in document.pages[0].raw_text


def test_pdf_scanned_candidate_is_explicit(tmp_path):
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with path.open("wb") as output:
        writer.write(output)
    document = parse_pdf(path)
    assert document.pages[0].parse_status == "ocr_candidate"
    assert "OCR" in (document.pages[0].warning or "")
    assert document.pages[0].confidence == 0.0


def test_docx_uses_section_locations_without_fake_page_numbers(tmp_path):
    path = tmp_path / "lesson.docx"
    document = Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("段落内容")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.save(path)
    parsed = parse_docx(path)
    assert parsed.pages[0].location_type == "section"
    assert "章节/段落" in (parsed.pages[0].location_label or "")
    assert parsed.pages[0].stable_location_key == "docx-section:1"
    assert any(block.block_type == "table" for block in parsed.pages[0].blocks)


def test_pdf_upload_reaches_persistent_queue(tmp_path, monkeypatch):
    from dataclasses import replace

    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    from app.core.config import settings
    from app.db.database import Base, get_db
    from app.services import deletion, uploads

    engine = create_engine(f"sqlite:///{(tmp_path / 'queue.sqlite3').as_posix()}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    monkeypatch.setattr(uploads, "settings", replace(settings, data_dir=tmp_path))
    monkeypatch.setattr(deletion, "_checkpoint", lambda: None)

    def override_db():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    try:
        client = TestClient(app)
        course = client.post("/api/v1/courses", json={"name": "PDF 课程"}).json()
        path = tmp_path / "empty.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        with path.open("wb") as output:
            writer.write(output)
        response = client.post(
            f"/api/v1/courses/{course['id']}/materials",
            files={"file": ("empty.pdf", path.read_bytes(), "application/pdf")},
            data={"lecture_title": "PDF 第一讲"},
        )
        assert response.status_code == 201, response.text
        assert response.json()["material"]["original_filename"] == "empty.pdf"
    finally:
        app.dependency_overrides.clear()
        engine.dispose()
