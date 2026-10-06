from __future__ import annotations

from dataclasses import replace
from io import BytesIO

from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker

from app.core.config import settings
from app.db.database import Base, get_db
from app.db.models import Course, KnowledgeNode, Note
from app.main import app
from app.services import deletion, uploads


def _pptx() -> BytesIO:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[5]).shapes.title.text = "知识点"
    result = BytesIO()
    presentation.save(result)
    result.seek(0)
    return result


def test_course_isolation_soft_delete_and_upload_queue(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{(tmp_path / 'test.sqlite3').as_posix()}")

    @event.listens_for(engine, "connect")
    def foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    monkeypatch.setattr(uploads, "settings", replace(settings, data_dir=tmp_path))
    # The endpoint test checks the reversible state transition; the checkpoint
    # filesystem round trip is tested separately.
    monkeypatch.setattr(deletion, "_checkpoint", lambda: None)

    def test_db():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = test_db
    try:
        client = TestClient(app)
        with factory() as db:
            a, b = Course(name="A"), Course(name="B")
            db.add_all([a, b])
            db.flush()
            node = KnowledgeNode(course_id=a.id, name="节点")
            db.add(node)
            db.flush()
            note = Note(course_id=a.id, knowledge_node_id=node.id, title="笔记", content_markdown="用户正文", user_locked=True)
            db.add(note)
            db.commit()
            a_id, b_id, note_id = a.id, b.id, note.id

        response = client.post(
            f"/api/v1/courses/{a_id}/materials",
            files={"file": ("one.pptx", _pptx(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            data={"lecture_title": "第一讲"},
            headers={"Idempotency-Key": "one"},
        )
        assert response.status_code == 201, response.text
        material_id = response.json()["material"]["id"]
        job_id = response.json()["job"]["id"]
        assert response.json()["job"]["status"] == "pending"
        assert client.get(f"/api/v1/jobs/{job_id}").json()["status"] == "pending"
        assert client.post(f"/api/v1/courses/{b_id}/assistant", json={
            "question": "测试", "material_id": material_id,
        }).status_code == 400
        assert client.get(f"/api/v1/materials/{material_id}/deletion-preview").json()["active_jobs"] == 1
        assert client.delete(f"/api/v1/materials/{material_id}").status_code == 409
        with factory() as db:
            from app.db.models import ProcessingJob
            job = db.get(ProcessingJob, job_id)
            job.status = "failed"
            db.commit()
        assert client.delete(f"/api/v1/materials/{material_id}").status_code == 200
        assert client.get(f"/api/v1/materials/{material_id}/pages").status_code == 404
        assert client.get(f"/api/v1/notes/{note_id}").json()["content_markdown"] == "用户正文"
        assert client.post(f"/api/v1/materials/{material_id}/restore").status_code == 200
        assert client.delete(f"/api/v1/courses/{a_id}").status_code == 200
        assert client.get(f"/api/v1/notes/{note_id}").status_code == 404
        assert client.post(f"/api/v1/courses/{a_id}/restore").status_code == 200
        assert client.get(f"/api/v1/notes/{note_id}").json()["content_markdown"] == "用户正文"
    finally:
        app.dependency_overrides.clear()
        engine.dispose()
