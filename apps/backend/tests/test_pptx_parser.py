"""PPTX 解析器的版面识别测试。

用例锁定的核心约束（正是本次「上传 PPTX 失败」的根因）：

1. 页脚/装饰带**不能**被当成页面标题。
   早期实现假定 `slide.shapes[0]` 就是标题，而示例课件的装饰文本框
   恰好排在正文之前，于是标题变成了
   `,   ·  ,   ·                   =                曾智 (北邮 AI 学院) …`。
2. 页脚文本**不能**混进 `raw_text`，否则会被喂给 LLM 当作正文。
3. 标题应取版面上沿的「页眉带」（章节名），而不是字号更大的正文块
   （正文 16pt > 页眉 14pt，仅按字号排序会误取正文）。
4. 整页只有页脚时应视为空页，而不是造出一个垃圾标题。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

from app.parsers.pptx_parser import parse_pptx

SLIDE_W = Emu(9144000)
SLIDE_H = Emu(3456051)


def _add_textbox(
    slide, text, *, top, left=Emu(0), width=Emu(4633595), height=Emu(400000), size=14.0
):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.text = text.split("\n")[0]
    for extra in text.split("\n")[1:]:
        frame.add_paragraph().text = extra
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
    return box


def _blank_deck() -> tuple[Presentation, object]:
    presentation = Presentation()
    presentation.slide_width = SLIDE_W
    presentation.slide_height = SLIDE_H
    blank = presentation.slide_layouts[6]
    return presentation, blank


def _save(presentation: Presentation, tmp_path: Path) -> Path:
    target = tmp_path / "deck.pptx"
    presentation.save(target)
    return target


FOOTER_TEXT = ",   ·  ,   ·                   =                \t曾智 (北邮 AI 学院)                                                              向量与空间                            2026 年 9 月 18  日        3 / 45"


def test_footer_shape_is_not_taken_as_title(tmp_path: Path) -> None:
    """页脚排在形状列表最前时，标题不应变成页脚文本。"""
    presentation, blank = _blank_deck()
    slide = presentation.slides.add_slide(blank)

    # 故意让页脚先添加——复现真实的形状顺序
    _add_textbox(slide, FOOTER_TEXT, top=Emu(3207283), height=Emu(248920), size=5.0)
    _add_textbox(
        slide,
        "1.1  向量空间    1.1.1 特征与向量\n特征与向量",
        top=Emu(-24548),
        height=Emu(462280),
        size=14.0,
    )
    _add_textbox(
        slide,
        "由单一数值构成的对待研究对象的量化评价，称作“标量”。",
        top=Emu(777538),
        height=Emu(2676525),
        size=16.0,
    )
    path = _save(presentation, tmp_path)

    page = parse_pptx(path)[0]

    assert page["title"] == "1.1  向量空间    1.1.1 特征与向量"
    assert "曾智" not in (page["title"] or "")
    assert "向量空间" in page["title"]


def test_footer_text_does_not_leak_into_raw_text(tmp_path: Path) -> None:
    """页脚与正文叠在同一 shape 时，页脚行要被逐行剔除。"""
    presentation, blank = _blank_deck()
    slide = presentation.slides.add_slide(blank)
    _add_textbox(
        slide,
        "待续\n\n\n\n\n\n\n\n\n\n\n\t,    ·   ,    ·                        =\n\n\t曾智 (北邮 AI 学院)                                                              向量与空间                            2026 年 9  月 18  日        45 / 45",
        top=Emu(1344496),
        height=Emu(2110104),
        size=24.0,
    )
    path = _save(presentation, tmp_path)

    page = parse_pptx(path)[0]

    assert "待续" in page["raw_text"]
    assert "曾智" not in page["raw_text"]
    assert "45 / 45" not in page["raw_text"]


def test_header_band_beats_larger_body_font(tmp_path: Path) -> None:
    """正文 16pt > 页眉 14pt，标题仍应取页眉带。"""
    presentation, blank = _blank_deck()
    slide = presentation.slides.add_slide(blank)
    _add_textbox(slide, FOOTER_TEXT, top=Emu(3207283), height=Emu(248920), size=5.0)
    _add_textbox(
        slide,
        "定理 1-1 Fn 中有如下的 n 个向量:\nεi = (0, · · · , 0, 1, 0, · · · , 0)",
        top=Emu(1255249),
        height=Emu(2199004),
        size=16.0,
    )
    _add_textbox(
        slide,
        "1.1  向量空间    1.1.3 基和维数\n基和维数",
        top=Emu(-24548),
        height=Emu(463550),
        size=14.0,
    )
    path = _save(presentation, tmp_path)

    page = parse_pptx(path)[0]

    assert page["title"] == "1.1  向量空间    1.1.3 基和维数"


def test_footer_only_page_yields_no_title(tmp_path: Path) -> None:
    """整页只有页脚时，标题应为 None，且 raw_text 为空。"""
    presentation, blank = _blank_deck()
    slide = presentation.slides.add_slide(blank)
    _add_textbox(slide, FOOTER_TEXT, top=Emu(3207283), height=Emu(248920), size=5.0)
    path = _save(presentation, tmp_path)

    page = parse_pptx(path)[0]

    assert page["title"] is None
    assert page["raw_text"] == ""
    assert page["blocks"] == []
    assert page["warning"] == "本页没有识别到文本内容"


def test_title_placeholder_wins_over_geometry(tmp_path: Path) -> None:
    """存在真正的标题占位符时，优先采用它。"""
    presentation, _blank = _blank_deck()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "欧式空间的定义"
    for paragraph in slide.shapes.title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(28)
    _add_textbox(
        slide,
        "正文段落，字号也不小，但不应覆盖标题占位符。",
        top=Emu(1500000),
        height=Emu(800000),
        size=20.0,
    )
    path = _save(presentation, tmp_path)

    page = parse_pptx(path)[0]

    assert page["title"] == "欧式空间的定义"


def test_multiline_header_returns_first_line(tmp_path: Path) -> None:
    """页眉带把章节名与小节名叠成两行时，标题取首行。"""
    presentation, blank = _blank_deck()
    slide = presentation.slides.add_slide(blank)
    _add_textbox(
        slide,
        "1.2 欧式空间    1.2.1  内积与投影\n内积与投影",
        top=Emu(-24548),
        height=Emu(463550),
        size=14.0,
    )
    path = _save(presentation, tmp_path)

    assert parse_pptx(path)[0]["title"] == "1.2 欧式空间    1.2.1  内积与投影"


def test_real_deck_no_page_has_footer_as_title() -> None:
    """用仓库里的真实课件回归：任何一页的标题都不应包含页脚特征文本。

    该文件正是触发上传失败的那份（`向量与空间_2026.pptx`）。缺失时跳过，
    以免在未携带样例数据的机器上误报。
    """
    sample = (
        Path(__file__).resolve().parents[3]
        / "data"
        / "originals"
        / "bc4287e2-9131-433e-bc88-aa2cd96d61f0.pptx"
    )
    if not sample.exists():
        pytest.skip("样例课件不在仓库中")

    pages = parse_pptx(sample)

    assert len(pages) == 45
    for page in pages:
        title = page["title"] or ""
        # 页脚带的特征串（页码、邮箱、讲者+课名+日期的连排）不能出现在标题里。
        # 注意：封面页正文本身就是作者信息，所以不禁止「曾智」单独出现，
        # 只禁止页脚**组合**特征——那才是解析错误的证据。
        assert "/ 45" not in title, f"第 {page['page_number']} 页标题混入页码"
        assert "zhi.zeng" not in title, f"第 {page['page_number']} 页标题混入邮箱"
        assert "北邮 AI 学院" not in title, f"第 {page['page_number']} 页标题混入页脚署名"
        assert not title.startswith(",   ·"), f"第 {page['page_number']} 页标题是装饰符"

    # 关键页的标题应落到章节名上，而不是装饰文本
    assert pages[2]["title"] == "1.1  向量空间    1.1.1 特征与向量"
    assert pages[23]["title"] == "1.2 欧式空间    1.2.1  内积与投影"
    # 纯装饰页不再产生垃圾标题
    assert pages[1]["title"] is None
